"""
Lê o relatório "Controle de Consultas" (PDF do CRESM) e gera a planilha para a campanha de
confirmação de consultas no WhatsApp.

Uso:
    python extrair_relatorio.py Relatorio.pdf
    python extrair_relatorio.py Relatorio.pdf --saida pasta_de_saida

Gera, na pasta de saída:
    campanha_AAAA-MM-DD.xlsx  abas "Campanha" (um paciente por linha) e "Pendências"
    campanha_AAAA-MM-DD.csv   mesma aba "Campanha" em CSV (separador ;)

Regras:
    - Um paciente com várias consultas no mesmo dia vira uma única linha; a coluna "consultas"
      traz uma linha por consulta: "• Psicologia – Ludmila de Souza Ribeiro".
    - Período pelo primeiro horário do dia: antes de 12h -> "a partir das 07h00 (matutino)",
      senão -> "a partir das 13h00 (vespertino)".
    - Telefone: o primeiro número válido do relatório; os demais vão para "telefones_reserva".
      Com a Evolution API configurada, usa o primeiro número que existe no WhatsApp, escrito como o
      WhatsApp o identifica (com ou sem o 9), para não abrir conversa duplicada.
    - Um mesmo telefone não é usado para dois pacientes: o segundo usa outro número dele ou vai para
      "Pendências".

Verificação no WhatsApp (opcional), por variáveis de ambiente:
    EVOLUTION_URL=http://localhost:8080 EVOLUTION_INSTANCIA=nome EVOLUTION_APIKEY=chave \
        python extrair_relatorio.py Relatorio.pdf

Dependências: pip install pdfplumber openpyxl
"""
import argparse
import csv
import os
import re
import sys
from collections import OrderedDict
from datetime import datetime

COLUNAS = ['prontuario', 'nome', 'telefone', 'telefones_reserva', 'whatsapp_verificado', 'data', 'periodo', 'consultas',
           'especialidades', 'profissionais', 'horarios', 'tipos_atendimento', 'municipio', 'idade']

# Posição horizontal (em pontos) das colunas da tabela de consultas, usada se o cabeçalho da
# tabela não for encontrado na página.
COLUNAS_PADRAO = {'horario': 95, 'prestador': 155, 'tipo': 393, 'especialidade': 467}

PARTICULAS = {'de', 'da', 'do', 'das', 'dos', 'e', 'di', 'du'}
RE_DATA = re.compile(r'^\d{2}/\d{2}/\d{4}$')
RE_HORA = re.compile(r'^\d{1,2}:\d{2}$')


def formatar_nome(texto):
    palavras = texto.lower().split()
    return ' '.join(p if i and p in PARTICULAS else p.capitalize() for i, p in enumerate(palavras))


def normalizar_telefone(texto):
    """Devolve o número no formato 55 + DDD + 9 dígitos, ou None se não for válido."""
    digitos = re.sub(r'\D', '', texto)
    if digitos.startswith('55') and len(digitos) in (12, 13):
        digitos = digitos[2:]
    if len(digitos) == 10 and digitos[2] in '6789':
        digitos = digitos[:2] + '9' + digitos[2:]  # celular antigo sem o 9
    if len(digitos) != 11 or digitos[2] != '9' or digitos[0] == '0':
        return None
    return '55' + digitos


def periodo(horario):
    hora = int(horario.split(':')[0])
    return 'a partir das 07h00 (matutino)' if hora < 12 else 'a partir das 13h00 (vespertino)'


def agrupar_linhas(palavras, tolerancia=3):
    linhas = []
    for palavra in sorted(palavras, key=lambda w: (round(w['top']), w['x0'])):
        if linhas and abs(linhas[-1][0]['top'] - palavra['top']) <= tolerancia:
            linhas[-1].append(palavra)
        else:
            linhas.append([palavra])
    return [sorted(linha, key=lambda w: w['x0']) for linha in linhas]


def texto(palavras):
    return ' '.join(w['text'] for w in palavras).strip()


def ler_pdf(caminho):
    """Lê o PDF e devolve (pacientes, totais_do_rodape, pendencias)."""
    import pdfplumber

    pacientes = []
    pendencias = []
    totais = {}
    atual = None
    ultima_consulta = None
    lendo_municipio = False
    colunas = dict(COLUNAS_PADRAO)

    with pdfplumber.open(caminho) as pdf:
        for numero_pagina, pagina in enumerate(pdf.pages, start=1):
            palavras = pagina.extract_words()
            # Ignora o cabeçalho da página (tudo até a linha "Período: ...")
            topo = next((w['bottom'] for w in palavras if w['text'] == 'Período:'), 0)
            linhas = agrupar_linhas([w for w in palavras if w['top'] > topo + 2])

            for linha in linhas:
                conteudo = texto(linha)
                primeira = linha[0]['text']

                if primeira == 'Total':
                    achado = re.match(r'Total de (Agendamentos|Pacientes):\s*(\d+)', conteudo)
                    if achado:
                        totais[achado.group(1).lower()] = int(achado.group(2))
                    continue

                if primeira == 'Prontuário:':
                    atual = {'prontuario': conteudo.split(':', 1)[1].strip(), 'nome': '',
                             'idade': '', 'telefones': [], 'municipio': '', 'consultas': [],
                             'pagina': numero_pagina}
                    pacientes.append(atual)
                    ultima_consulta = None
                    lendo_municipio = False
                    continue

                if atual is None:
                    continue

                if primeira == 'Paciente:':
                    atual['nome'] = conteudo.split(':', 1)[1].strip()
                    continue

                if primeira == 'Idade:':
                    rotulo_municipio = next((w for w in linha if w['text'] == 'Município:'), None)
                    limite = rotulo_municipio['x0'] if rotulo_municipio else float('inf')
                    antes = texto([w for w in linha if w['x0'] < limite])
                    achado = re.match(r'Idade:\s*(\d*)\s*Telefone Paciente:\s*(.*)', antes)
                    if achado:
                        atual['idade'] = achado.group(1)
                        atual['telefones'] = [t.strip() for t in achado.group(2).split('/') if t.strip()]
                    if rotulo_municipio:
                        atual['municipio'] = texto([w for w in linha if w['x0'] > limite])
                    lendo_municipio = True
                    continue

                # Cabeçalho da tabela de consultas: atualiza a posição das colunas
                if primeira in ('Data', 'Consulta'):
                    lendo_municipio = False
                    for w in linha:
                        if w['text'] == 'Agendamento':
                            colunas['horario'] = w['x0'] - 5
                        elif w['text'] == 'Prestador':
                            colunas['prestador'] = w['x0'] - 5
                        elif w['text'] in ('Atendimento', 'Tipo'):
                            colunas['tipo'] = w['x0'] - 5
                        elif w['text'] == 'Especialidade':
                            colunas['especialidade'] = w['x0'] - 5
                    continue

                if lendo_municipio:  # município quebrado em mais de uma linha
                    atual['municipio'] = f"{atual['municipio']} {conteudo}".strip()
                    continue

                partes = {'data': [], 'horario': [], 'prestador': [], 'tipo': [], 'especialidade': []}
                for w in linha:
                    if w['x0'] >= colunas['especialidade']:
                        partes['especialidade'].append(w)
                    elif w['x0'] >= colunas['tipo']:
                        partes['tipo'].append(w)
                    elif w['x0'] >= colunas['prestador']:
                        partes['prestador'].append(w)
                    elif w['x0'] >= colunas['horario']:
                        partes['horario'].append(w)
                    else:
                        partes['data'].append(w)
                partes = {k: texto(v) for k, v in partes.items()}

                if RE_DATA.match(partes['data']) and RE_HORA.match(partes['horario']):
                    ultima_consulta = partes
                    atual['consultas'].append(ultima_consulta)
                elif ultima_consulta is not None and not partes['data'] and not partes['horario']:
                    # continuação de célula quebrada (ex.: "PRIMEIRA" / "CONSULTA")
                    for campo in ('prestador', 'tipo', 'especialidade'):
                        if partes[campo]:
                            ultima_consulta[campo] = f'{ultima_consulta[campo]} {partes[campo]}'.strip()
                else:
                    pendencias.append({'prontuario': atual['prontuario'], 'nome': atual['nome'],
                                       'motivo': f'linha não reconhecida (página {numero_pagina}): {conteudo}'})

    return pacientes, totais, pendencias


def chave_telefone(numero):
    """DDD + 8 últimos dígitos: identifica o mesmo telefone com ou sem o 9."""
    digitos = re.sub(r'\D', '', numero)
    if digitos.startswith('55') and len(digitos) in (12, 13):
        digitos = digitos[2:]
    return digitos[:2] + digitos[-8:]


def config_evolution():
    url = os.environ.get('EVOLUTION_URL', '').rstrip('/')
    instancia = os.environ.get('EVOLUTION_INSTANCIA', '')
    apikey = os.environ.get('EVOLUTION_APIKEY', '')
    if url and instancia and apikey:
        return {'url': url, 'instancia': instancia, 'apikey': apikey}
    return None


def verificar_whatsapp(numeros, config, tamanho_lote=50):
    """Consulta a Evolution API e devolve {numero: numero_no_whatsapp ou None}.

    O número devolvido é o identificador real da conta (com ou sem o 9), que é o que evita
    abrir uma segunda conversa quando o paciente responde.
    """
    import json
    import urllib.request
    from urllib.parse import quote

    endereco = f"{config['url']}/chat/whatsappNumbers/{quote(config['instancia'])}"
    resultado = {}
    numeros = list(dict.fromkeys(numeros))
    for inicio in range(0, len(numeros), tamanho_lote):
        lote = numeros[inicio:inicio + tamanho_lote]
        requisicao = urllib.request.Request(
            endereco, data=json.dumps({'numbers': lote}).encode(), method='POST',
            headers={'Content-Type': 'application/json', 'apikey': config['apikey']})
        with urllib.request.urlopen(requisicao, timeout=60) as resposta:
            itens = json.load(resposta)
        por_chave = {}
        for item in itens:
            jid = item.get('jid') or ''
            numero_real = jid.split('@')[0] if item.get('exists') and jid else None
            for campo in (item.get('number') or '', jid.split('@')[0]):
                if campo:
                    por_chave.setdefault(chave_telefone(campo), numero_real)
        for numero in lote:
            resultado[numero] = por_chave.get(chave_telefone(numero))
    return resultado


def montar_campanha(pacientes, pendencias, whatsapp=None):
    """whatsapp: resultado de verificar_whatsapp(), ou None para não verificar."""
    linhas = []
    usados = {}  # chave_telefone -> (prontuario, nome) do paciente que já usa o número
    for paciente in pacientes:
        if not paciente['consultas']:
            pendencias.append({'prontuario': paciente['prontuario'], 'nome': paciente['nome'],
                               'motivo': 'paciente sem consultas no relatório'})
            continue

        validos = []
        for numero in paciente['telefones']:
            normalizado = normalizar_telefone(numero)
            if normalizado and chave_telefone(normalizado) not in {chave_telefone(v) for v in validos}:
                validos.append(normalizado)
        nome = formatar_nome(paciente['nome'])
        if not validos:
            pendencias.append({'prontuario': paciente['prontuario'], 'nome': nome,
                               'motivo': f"sem telefone válido: {' / '.join(paciente['telefones']) or 'vazio'}"})
            continue

        if whatsapp is not None:
            candidatos = [whatsapp.get(v) for v in validos if whatsapp.get(v)]
            if not candidatos:
                pendencias.append({'prontuario': paciente['prontuario'], 'nome': nome,
                                   'motivo': f"nenhum telefone tem WhatsApp: {' / '.join(validos)}"})
                continue
        else:
            candidatos = validos

        livres = [c for c in candidatos
                  if usados.get(chave_telefone(c), (paciente['prontuario'],))[0] == paciente['prontuario']]
        if not livres:
            dono = usados[chave_telefone(candidatos[0])]
            pendencias.append({'prontuario': paciente['prontuario'], 'nome': nome,
                               'motivo': f'telefone compartilhado com {dono[1]} (prontuário {dono[0]}) '
                                         'e sem outro número disponível'})
            continue
        telefone = livres[0]
        usados[chave_telefone(telefone)] = (paciente['prontuario'], nome)
        reserva = [c for c in candidatos if c != telefone]

        por_data = OrderedDict()
        for consulta in paciente['consultas']:
            por_data.setdefault(consulta['data'], []).append(consulta)

        for data, consultas in por_data.items():
            consultas.sort(key=lambda c: datetime.strptime(c['horario'], '%H:%M'))
            linhas.append({
                'prontuario': paciente['prontuario'],
                'nome': nome,
                'telefone': telefone,
                'telefones_reserva': ' / '.join(reserva),
                'whatsapp_verificado': 'sim' if whatsapp is not None else 'não',
                'data': data,
                'periodo': periodo(consultas[0]['horario']),
                'consultas': '\n'.join(f"• {c['especialidade'].capitalize()} – {formatar_nome(c['prestador'])}"
                                       for c in consultas),
                'especialidades': ', '.join(c['especialidade'].capitalize() for c in consultas),
                'profissionais': ', '.join(formatar_nome(c['prestador']) for c in consultas),
                'horarios': ', '.join(c['horario'] for c in consultas),
                'tipos_atendimento': ', '.join(c['tipo'].capitalize() for c in consultas),
                'municipio': paciente['municipio'],
                'idade': paciente['idade'],
            })

    # Reserva não pode ter número que ficou com outro paciente (evita mensagem no telefone errado)
    for linha in linhas:
        linha['telefones_reserva'] = ' / '.join(
            n for n in linha['telefones_reserva'].split(' / ')
            if n and usados.get(chave_telefone(n), (linha['prontuario'],))[0] == linha['prontuario'])
    return linhas


def mensagem_exemplo(linha):
    return (
        'Olá!\nEsta é uma mensagem automática do Ambulatório do CRESM para confirmação da sua consulta.\n'
        f"🗓️ Data: {linha['data']} {linha['periodo']}.\n"
        'O atendimento será realizado por ordem de chegada.\n'
        f"👤 Nome: {linha['nome']}\n"
        f"🗂️ Prontuário: {linha['prontuario']}\n"
        f"🩺 Consultas:\n{linha['consultas']}\n"
        '📍 Endereço do CRESM: ...'
    )


def salvar(linhas, pendencias, pasta):
    import openpyxl
    from openpyxl.styles import Alignment, Font

    datas = sorted({l['data'] for l in linhas}, key=lambda d: datetime.strptime(d, '%d/%m/%Y'))
    sufixo = datetime.strptime(datas[0], '%d/%m/%Y').strftime('%Y-%m-%d') if datas else 'sem-data'
    if len(datas) > 1:
        sufixo += '_a_' + datetime.strptime(datas[-1], '%d/%m/%Y').strftime('%Y-%m-%d')
    os.makedirs(pasta, exist_ok=True)
    base = os.path.join(pasta, f'campanha_{sufixo}')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Campanha'
    ws.append(COLUNAS)
    for linha in linhas:
        ws.append([linha[c] for c in COLUNAS])
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for celula in ws['H'][1:]:
        celula.alignment = Alignment(wrap_text=True, vertical='top')
    larguras = {'A': 12, 'B': 38, 'C': 16, 'D': 32, 'E': 12, 'F': 12, 'G': 32, 'H': 60, 'I': 30, 'J': 50}
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura

    wp = wb.create_sheet('Pendências')
    wp.append(['prontuario', 'nome', 'motivo'])
    for p in pendencias:
        wp.append([p['prontuario'], p['nome'], p['motivo']])
    for celula in wp[1]:
        celula.font = Font(bold=True)
    wp.column_dimensions['B'].width = 38
    wp.column_dimensions['C'].width = 80

    wb.save(base + '.xlsx')
    with open(base + '.csv', 'w', newline='', encoding='utf-8-sig') as arquivo:
        escritor = csv.DictWriter(arquivo, fieldnames=COLUNAS, delimiter=';')
        escritor.writeheader()
        escritor.writerows(linhas)
    return base + '.xlsx', base + '.csv'


def main():
    parser = argparse.ArgumentParser(description='Gera a planilha da campanha a partir do relatório PDF')
    parser.add_argument('pdf', help='relatório "Controle de Consultas" em PDF')
    parser.add_argument('--saida', default='.', help='pasta onde salvar a planilha (padrão: atual)')
    args = parser.parse_args()

    pacientes, totais, pendencias = ler_pdf(args.pdf)
    agendamentos = sum(len(p['consultas']) for p in pacientes)

    whatsapp = None
    config = config_evolution()
    if config:
        numeros = [n for p in pacientes for n in map(normalizar_telefone, p['telefones']) if n]
        print(f'Verificando {len(set(numeros))} números na Evolution API ({config["url"]})...')
        try:
            whatsapp = verificar_whatsapp(numeros, config)
        except Exception as erro:
            print(f'ERRO ao consultar a Evolution API: {type(erro).__name__}: {erro}')
            print('Planilha NÃO gerada. Confira EVOLUTION_URL, EVOLUTION_INSTANCIA e EVOLUTION_APIKEY.')
            return 2
    linhas = montar_campanha(pacientes, pendencias, whatsapp)
    xlsx, csv_ = salvar(linhas, pendencias, args.saida)

    print(f'Pacientes lidos:     {len(pacientes)} (relatório: {totais.get("pacientes", "?")})')
    print(f'Agendamentos lidos:  {agendamentos} (relatório: {totais.get("agendamentos", "?")})')
    print(f'Linhas na campanha:  {len(linhas)} '
          f'({sum(1 for l in linhas if chr(10) in l["consultas"])} com mais de uma consulta)')
    print(f'Pendências:          {len(pendencias)}')
    print(f'Planilha: {os.path.abspath(xlsx)}')
    print(f'CSV:      {os.path.abspath(csv_)}')

    if whatsapp is None:
        print('\nAviso: números NÃO verificados no WhatsApp (sem EVOLUTION_URL/EVOLUTION_INSTANCIA/'
              'EVOLUTION_APIKEY). Pode haver conversa duplicada por causa do 9º dígito.')

    divergente = (totais.get('pacientes') not in (None, len(pacientes))
                  or totais.get('agendamentos') not in (None, agendamentos))
    if divergente or not totais:
        print('\n*** ATENÇÃO: os totais lidos não batem com o rodapé do relatório. '
              'Confira a planilha antes de enviar a campanha. ***')

    exemplo = next((l for l in linhas if '\n' in l['consultas']), linhas[0] if linhas else None)
    if exemplo:
        print('\nExemplo de mensagem:\n' + mensagem_exemplo(exemplo))
    return 1 if divergente else 0


if __name__ == '__main__':
    sys.exit(main())
